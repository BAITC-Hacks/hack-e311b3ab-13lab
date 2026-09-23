from io import BytesIO

from docx import Document


def markdown(meeting):
    analysis = meeting["analysis"]
    lines = [f"# {meeting['title']}", "", f"Дата совещания: {meeting['meeting_date']}", "", "Черновик ИИ. Требуется проверка секретарём.", "", "## Саммари", analysis["summary"], "", "## Решения"]
    lines.extend(f"- {decision}" for decision in analysis["decisions"])
    lines += ["", "## Поручения"]
    for index, action in enumerate(analysis["actions"], 1):
        lines.extend([f"### {index}. {action['title']}", f"Ответственный: {action['owner'] or 'Не указан'}", f"Срок: {action['due_date'] or action['deadline_text'] or 'Не указан'}", f"Статус: {action['status']}", f"Цитата: {action['evidence']}", ""])
    lines += ["## Замечания"] + [f"- {warning}" for warning in analysis["warnings"]]
    lines += ["", "## Транскрипт", meeting["transcript"]]
    return "\n".join(lines)


def docx(meeting):
    document = Document()
    document.add_heading(meeting["title"], 0)
    document.add_paragraph(f"Дата совещания: {meeting['meeting_date']}")
    document.add_paragraph("Черновик ИИ. Требуется проверка секретарём.")
    analysis = meeting["analysis"]
    document.add_heading("Саммари", 1)
    document.add_paragraph(analysis["summary"])
    document.add_heading("Решения", 1)
    for decision in analysis["decisions"]:
        document.add_paragraph(decision, style="List Bullet")
    document.add_heading("Поручения", 1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Срок"]):
        cell.text = label
    for action in analysis["actions"]:
        cells = table.add_row().cells
        cells[0].text = action["title"]
        cells[1].text = action["owner"] or "Не указан"
        cells[2].text = action["due_date"] or action["deadline_text"] or "Не указан"
    document.add_heading("Основания и проверка", 1)
    for index, action in enumerate(analysis["actions"], 1):
        document.add_paragraph(f"{index}. {action['evidence']}")
        document.add_paragraph(f"Статус: {action['status']}; проверено: {'нет' if action['needs_review'] else 'да'}")
    for warning in analysis["warnings"]:
        document.add_paragraph(warning)
    document.add_heading("Транскрипт", 1)
    for segment in meeting["segments"]:
        speaker = meeting.get("speaker_names", {}).get(segment["speaker"], segment["speaker"]) or "Говорящий не определён"
        timestamp = f"[{segment['start']:.1f} с] " if segment["start"] is not None else ""
        document.add_paragraph(f"{timestamp}{speaker}: {segment['text']}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
