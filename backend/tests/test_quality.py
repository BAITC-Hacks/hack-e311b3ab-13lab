import json
import unittest
from io import BytesIO

from docx import Document

from app.config import Settings
from app.deadlines import ground_deadlines
from app.exports import docx, markdown
from app.models import Action, Analysis, DeadlineAlternative, Segment
from app.pipeline import Provider, validate_evidence
from app.quality import numeric_fragments, prefer_summary, recover_explicit_owner


class QualityTests(unittest.TestCase):
    def test_recover_direct_addressee_without_claiming_audio_confirmation(self):
        for quote in ("Нурлан Сагатович, проведите инструктаж.", "Нурлан Сагатович, ну-ка, подскажите нам всем, когда был инструктаж, вот и проводите, ждем от вас на следующей неделе инструктаж."):
            action = Action(title="Инструктаж", evidence=quote, segment_ids=["s1"])
            recover_explicit_owner(action, [Segment(id="s1", text=quote)])
            self.assertEqual(action.owner, "Нурлан Сагатович")
            self.assertEqual(action.owner_evidence, quote)
            self.assertTrue(action.owner_uncertain)
            self.assertTrue(action.needs_review)
            self.assertIn("по аудио", action.review_questions[0])

    def test_do_not_infer_owner_from_request_to_speak_or_damaged_name(self):
        for quote in ("Асхат Ерланович, можно добавить? Свяжитесь с Нурланом.", "Нурлан Сагатовича, свяжитесь с Нурланом.", "Солгатович, соберите совещание.", "Нурлан Сагатович, проведите инструктаж. Тимур Болатович, подготовьте отчёт."):
            action = Action(title="Поручение", evidence=quote, segment_ids=["s1"])
            recover_explicit_owner(action, [Segment(id="s1", text=quote)])
            self.assertIsNone(action.owner)

    def test_owner_recovery_requires_own_source_and_preserves_existing_owner(self):
        quote = "Нурлан Сагатович, проведите инструктаж."
        action = Action(title="Инструктаж", evidence="Проведите инструктаж.", segment_ids=["s1"])
        recover_explicit_owner(action, [Segment(id="s1", text=quote)])
        self.assertIsNone(action.owner)
        action.evidence = quote
        action.owner = "Подтверждённый исполнитель"
        recover_explicit_owner(action, [Segment(id="s1", text=quote)])
        self.assertEqual(action.owner, "Подтверждённый исполнитель")

    def test_summary_repair_must_keep_every_previously_covered_fragment(self):
        segments = [Segment(id="s1", text="Загрузка 71%. Потери 8%. Рост 60%.")]
        first = "Загрузка 71%."
        self.assertEqual(prefer_summary(first, "Потери 8%. Рост 60%.", segments), first)
        self.assertEqual(prefer_summary(first, "Другой пересказ. Загрузка 71%.", segments), first)
        improved = "Загрузка 71%. Потери 8%."
        self.assertEqual(prefer_summary(first, improved, segments), improved)
        self.assertEqual(prefer_summary(first, improved + " [Цитата не совпала с источником — требуется проверка]", segments), first)

    def test_numbers_require_source_context_not_just_digits(self):
        source = Segment(id="s1", text="Загрузка семьдесят один процент. Потери 8%.")
        fragments = numeric_fragments([source], "Готовность 71%. Потери 8%.")
        self.assertEqual([item.included for item in fragments], [False, True])
        self.assertEqual(numeric_fragments([Segment(id="s2", text="Жеті күн. Перенос на полгода.")], "")[1].text, "Перенос на полгода.")

    def test_details_require_exact_evidence(self):
        action = Action(title="Отчёт", evidence="Доложите лично.", deliverable="Личный доклад", deliverable_evidence="Доложите лично.", condition="Расторгнуть договор", condition_evidence="Выдуманная цитата", issued_by="Асхат")
        result = validate_evidence(Analysis(summary="", actions=[action]), [Segment(id="s1", text=action.evidence)])
        self.assertEqual(result.actions[0].deliverable, "Личный доклад")
        self.assertIsNone(result.actions[0].condition)
        self.assertIsNone(result.actions[0].issued_by)

    def test_chair_can_be_owner_with_evidence(self):
        quote = "Асхат, подготовьте отчёт."
        action = Action(title="Отчёт", owner="Асхат", evidence=quote, owner_evidence=quote, issued_by="Тимур", issued_by_evidence="Тимур поручил отчёт.")
        result = validate_evidence(Analysis(summary="", actions=[action]), [Segment(id="s1", text=quote + " Тимур поручил отчёт.")])
        self.assertEqual(result.actions[0].owner, "Асхат")
        self.assertEqual(result.actions[0].issued_by, "Тимур")

    def test_event_deadline_keeps_text_without_date(self):
        action = Action(title="Справка", evidence="Подготовьте справку после совещания с подрядчиками.", deadline_text="после совещания с подрядчиками", due_date="2026-09-25")
        result = ground_deadlines(Analysis(summary="", actions=[action]), "2026-09-23")
        self.assertEqual(result.actions[0].deadline_resolution, "event")
        self.assertIsNone(result.actions[0].due_date)
        self.assertEqual(result.actions[0].deadline_text, "после совещания с подрядчиками")

    def test_conflicts_do_not_silently_pick_date(self):
        source = Segment(id="s1", text="Срок две недели или максимум десять дней, уточним.")
        action = Action(title="Отчёт", evidence=source.text, deadline_text="максимум десять дней", deadline_alternatives=[DeadlineAlternative(text="две недели", segment_id="s1"), DeadlineAlternative(text="максимум десять дней", segment_id="s1"), DeadlineAlternative(text="завтра", segment_id="invented")])
        result = ground_deadlines(validate_evidence(Analysis(summary="", actions=[action]), [source]), "2026-09-23")
        self.assertEqual(result.actions[0].deadline_resolution, "conflict")
        self.assertIsNone(result.actions[0].due_date)
        self.assertEqual(len(result.actions[0].deadline_alternatives), 2)

    def test_explicit_final_deadline_resolves(self):
        action = Action(title="Отчёт", evidence="Согласовали максимум десять дней.", deadline_text="максимум десять дней")
        result = ground_deadlines(Analysis(summary="", actions=[action]), "2026-09-23")
        self.assertEqual(str(result.actions[0].due_date), "2026-10-03")
        self.assertEqual(result.actions[0].deadline_resolution, "resolved")

    def test_exports_keep_details_and_conflict_sources(self):
        action = Action(title="Отчёт", evidence="Доложите лично.", deliverable="Личный доклад", condition="При нарушениях", deadline_resolution="conflict", deadline_alternatives=[DeadlineAlternative(text="две недели", segment_id="s1")])
        meeting = {"title":"Тест", "meeting_date":"2026-09-23", "analysis":Analysis(summary="Тест", actions=[action]).model_dump(mode="json"), "transcript":"Доложите лично.", "segments":[]}
        document = Document(BytesIO(docx(meeting)))
        self.assertIn("Личный доклад", document.tables[0].rows[1].cells[0].text)
        self.assertIn("две недели [s1]", document.tables[0].rows[1].cells[2].text)
        self.assertIn("При нарушениях", markdown(meeting))


class SummaryTests(unittest.IsolatedAsyncioTestCase):
    async def test_regressive_corrective_pass_keeps_first_summary(self):
        provider = Provider(Settings.from_env())
        responses = iter(["Загрузка 71%.", "Потери 8%."])

        async def request(endpoint, **kwargs):
            return {"choices": [{"message": {"content": json.dumps({"summary": next(responses)})}}]}

        provider.request = request
        result = await provider.summarize([Segment(id="s1", text="Загрузка 71%. Потери 8%.")])
        self.assertEqual(result, "Загрузка 71%.")

    async def test_one_corrective_pass_preserves_numeric_context(self):
        provider = Provider(Settings.from_env())
        calls = []

        async def request(endpoint, **kwargs):
            calls.append(json.loads(json.dumps(kwargs)))
            summary = "Обсудили загрузку." if len(calls) == 1 else "Загрузка 71%."
            return {"choices":[{"finish_reason":"stop", "message":{"content":json.dumps({"summary":summary})}}]}

        provider.request = request
        result = await provider.summarize([Segment(id="s1", text="Загрузка 71%.")])
        self.assertEqual(result, "Загрузка 71%.")
        self.assertEqual(len(calls), 2)
        self.assertIn("Загрузка 71%.", calls[1]["json"]["messages"][-1]["content"])

    async def test_correction_is_bounded_when_model_omits_numbers(self):
        provider = Provider(Settings.from_env())
        calls = []

        async def request(endpoint, **kwargs):
            calls.append(endpoint)
            return {"choices":[{"finish_reason":"stop", "message":{"content":'{"summary":"Обсудили загрузку."}'}}]}

        provider.request = request
        result = await provider.summarize([Segment(id="s1", text="Загрузка 71%.")])
        self.assertEqual(len(calls), 2)
        self.assertFalse(numeric_fragments([Segment(id="s1", text="Загрузка 71%.")], result)[0].included)
