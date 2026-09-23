from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from app.source_review import review_lines

STATUS_LABELS = {"open": "Открыто", "in_progress": "В работе", "done": "Выполнено"}


def _deadline(action):
    value = action["due_date"] or action["deadline_text"] or "Не указан"
    alternatives = [f"{item['text']} [{item['segment_id']}]" for item in action.get("deadline_alternatives", [])]
    return str(value) + ("; варианты: " + "; ".join(alternatives) if alternatives else "")


def _details(action):
    labels = {"issued_by": "Поручил(а)", "deliverable": "Результат", "condition": "Условие", "owner_evidence": "Основание исполнителя", "deliverable_evidence": "Основание результата", "condition_evidence": "Основание условия"}
    return [f"{label}: {action[field]}" for field, label in labels.items() if action.get(field)] + action.get("review_questions", [])


def _owner(action, people):
    owner = action["owner"] or "Не указан"
    assignee = people.get(action.get("assignee_id") or "")
    return f"{owner} → {assignee}" if assignee and assignee != action["owner"] else (assignee or owner)


def _approval_line(meeting, people):
    if meeting.get("status") == "approved" and meeting.get("approved_at"):
        return f"Протокол утверждён {meeting['approved_at'][:10]}: {people.get(meeting.get('approved_by') or '', 'пользователь удалён')}."
    return "Черновик ИИ. Требуется проверка секретарём."


def markdown(meeting, people=None):
    people = people or {}
    analysis = meeting["analysis"]
    lines = [f"# {meeting['title']}", "", f"Дата совещания: {meeting['meeting_date']}", "", _approval_line(meeting, people), ""]
    if meeting.get("chair_id") in people:
        lines += [f"Председатель: {people[meeting['chair_id']]}", ""]
    participants = [people[user_id] for user_id in meeting.get("participant_ids", []) if user_id in people]
    if participants:
        lines += [f"Участники: {', '.join(participants)}", ""]
    lines += ["## Саммари", analysis["summary"], "", "## Решения"]
    lines.extend(f"- {decision}" for decision in analysis["decisions"])
    lines += ["", "## Поручения"]
    for index, action in enumerate(analysis["actions"], 1):
        lines.extend([f"### {index}. {action['title']}", f"Ответственный: {_owner(action, people)}", f"Срок: {_deadline(action)}", f"Статус: {STATUS_LABELS.get(action['status'], action['status'])}", f"Цитата: {action['evidence']}", ""])
        lines.extend(_details(action))
    lines += ["## Замечания"] + [f"- {warning}" for warning in analysis["warnings"]]
    if review_lines(meeting):
        lines += ["", "## Ручная сверка с аудио (исходный ASR сохранён)", *review_lines(meeting)]
    lines += ["", "## Транскрипт", meeting["transcript"]]
    return "\n".join(lines)


def docx(meeting, people=None):
    people = people or {}
    document = Document()
    title = document.add_heading(meeting["title"], 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.underline = False
    for properties in (title._p.get_or_add_pPr(), title.style.element.get_or_add_pPr()):
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
    document.add_paragraph(f"Дата совещания: {meeting['meeting_date']}")
    document.add_paragraph(_approval_line(meeting, people))
    if meeting.get("chair_id") in people:
        document.add_paragraph(f"Председатель: {people[meeting['chair_id']]}")
    participants = [people[user_id] for user_id in meeting.get("participant_ids", []) if user_id in people]
    if participants:
        document.add_paragraph(f"Участники: {', '.join(participants)}")
    analysis = meeting["analysis"]
    document.add_heading("Саммари", 1)
    document.add_paragraph(analysis["summary"])
    document.add_heading("Решения", 1)
    for decision in analysis["decisions"]:
        document.add_paragraph(decision, style="List Bullet")
    document.add_heading("Поручения", 1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, (2.7, 1.3, 1.3, 0.7)):
        column.width = Inches(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, label in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Срок", "Статус"]):
        cell.text = label
    for action in analysis["actions"]:
        cells = table.add_row().cells
        cells[0].text = "\n".join([action["title"], *[f"{label}: {action[field]}" for field, label in (("deliverable", "Результат"), ("condition", "Условие")) if action.get(field)]])
        cells[1].text = _owner(action, people)
        cells[2].text = _deadline(action)
        cells[3].text = STATUS_LABELS.get(action["status"], action["status"])
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, (2.7, 1.3, 1.3, 0.7)):
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_heading("Основания и проверка", 1)
    for index, action in enumerate(analysis["actions"], 1):
        document.add_paragraph(f"{index}. {action['evidence']}")
        document.add_paragraph(f"Проверено: {'нет' if action['needs_review'] else 'да'}")
        for detail in _details(action):
            document.add_paragraph(detail)
    for warning in analysis["warnings"]:
        document.add_paragraph(warning)
    if review_lines(meeting):
        document.add_heading("Ручная сверка с аудио (исходный ASR сохранён)", 1)
        for line in review_lines(meeting):
            document.add_paragraph(line)
    document.add_heading("Транскрипт", 1)
    for segment in meeting["segments"]:
        speaker = meeting.get("speaker_names", {}).get(segment["speaker"], segment["speaker"]) or "Говорящий не определён"
        timestamp = f"[{segment['start']:.1f} с] " if segment["start"] is not None else ""
        document.add_paragraph(f"{timestamp}{speaker}: {segment['text']}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", None),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]


def find_fonts(settings):
    """Return (regular, bold) TTF paths with Cyrillic glyphs, or None when PDF is unavailable."""
    candidates = [(settings.pdf_font_path, settings.pdf_font_bold_path or None)] if settings.pdf_font_path else FONT_CANDIDATES
    for regular, bold in candidates:
        if regular and Path(regular).is_file():
            return regular, bold if bold and Path(bold).is_file() else regular
    return None


def pdf(meeting, fonts, people=None):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    people = people or {}
    regular, bold = fonts
    document = FPDF()
    document.set_auto_page_break(True, margin=15)
    document.add_font("Main", "", regular)
    document.add_font("Main", "B", bold)
    document.set_title(meeting["title"])
    document.add_page()

    def paragraph(value, size=10, style="", gap=1.5):
        document.set_font("Main", style, size)
        document.multi_cell(0, size * 0.5, value, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        document.ln(gap)

    def heading(value):
        document.ln(2)
        paragraph(value, 13, "B", 1)

    analysis = meeting["analysis"]
    paragraph(meeting["title"], 18, "B", 2)
    paragraph(f"Дата совещания: {meeting['meeting_date']}")
    paragraph(_approval_line(meeting, people))
    if meeting.get("chair_id") in people:
        paragraph(f"Председатель: {people[meeting['chair_id']]}")
    participants = [people[user_id] for user_id in meeting.get("participant_ids", []) if user_id in people]
    if participants:
        paragraph(f"Участники: {', '.join(participants)}")
    heading("Саммари")
    paragraph(analysis["summary"] or "—")
    heading("Решения")
    for decision in analysis["decisions"] or ["—"]:
        paragraph(f"• {decision}")
    heading("Поручения")
    if analysis["actions"]:
        document.set_font("Main", "", 9)
        with document.table(col_widths=(8, 44, 22, 14, 12), text_align="LEFT", line_height=4.5) as table:
            header = table.row()
            for label in ("№", "Поручение", "Ответственный", "Срок", "Статус"):
                header.cell(label)
            for index, action in enumerate(analysis["actions"], 1):
                row = table.row()
                for value in (str(index), action["title"], _owner(action, people), _deadline(action), STATUS_LABELS.get(action["status"], action["status"])):
                    row.cell(value)
        heading("Основания")
        for index, action in enumerate(analysis["actions"], 1):
            paragraph(f"{index}. «{action['evidence']}» Проверено: {'нет' if action['needs_review'] else 'да'}.", 9)
            for detail in _details(action):
                paragraph(detail, 9)
    else:
        paragraph("Поручения не выделены.")
    if analysis["warnings"]:
        heading("Замечания")
        for warning in analysis["warnings"]:
            paragraph(f"• {warning}", 9)
    if review_lines(meeting):
        heading("Ручная сверка с аудио (исходный ASR сохранён)")
        for line in review_lines(meeting):
            paragraph(line, 9)
    heading("Транскрипт")
    for segment in meeting["segments"]:
        speaker = meeting.get("speaker_names", {}).get(segment["speaker"], segment["speaker"]) or "Говорящий не определён"
        timestamp = f"[{segment['start']:.1f} с] " if segment["start"] is not None else ""
        paragraph(f"{timestamp}{speaker}: {segment['text']}", 9, gap=0.8)
    return bytes(document.output())
